from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any
from uuid import UUID

import xlrd
from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.database import get_db
from app.dataset_analytics import ALIASES, dashboard_analysis, forecast_analysis
from app.models import ImportedDataset, User
from app.security import get_current_user, log_activity

router = APIRouter(prefix="/datasets", tags=["Dataset Import"])

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 200
SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xls"}


def _header(value: Any, index: int) -> str:
    text = str(value or f"column_{index + 1}").strip().lower()
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    return text or f"column_{index + 1}"


def _unique_headers(values: list[Any]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for index, value in enumerate(values):
        base = _header(value, index)
        seen[base] = seen.get(base, 0) + 1
        result.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return result


def _json_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _number_or_text(value: str) -> Any:
    text = value.strip()
    if not text:
        return None
    lowered = text.lower()
    if lowered in {"true", "yes"}:
        return True
    if lowered in {"false", "no"}:
        return False
    cleaned = text.replace(",", "").replace("₹", "").replace("$", "")
    try:
        return float(cleaned) if "." in cleaned else int(cleaned)
    except ValueError:
        return text


def _rows_from_matrix(matrix: list[list[Any]], sheet_name: str | None) -> tuple[list[str], list[dict[str, Any]], str | None]:
    matrix = [list(row[:MAX_COLUMNS]) for row in matrix if any(cell not in (None, "") for cell in row)]
    if not matrix:
        raise HTTPException(status_code=422, detail="The uploaded file does not contain any records.")
    candidates = matrix[:20]
    header_index = max(range(len(candidates)), key=lambda index: sum(cell not in (None, "") for cell in candidates[index]))
    columns = _unique_headers(candidates[header_index])
    rows = []
    for raw in matrix[header_index + 1:MAX_ROWS + header_index + 1]:
        padded = list(raw) + [None] * (len(columns) - len(raw))
        if not any(value not in (None, "") for value in padded[:len(columns)]):
            continue
        rows.append({column: _json_value(padded[index]) for index, column in enumerate(columns)})
    if not rows:
        raise HTTPException(status_code=422, detail="A header was found, but the file has no data rows.")
    return columns, rows, sheet_name


def _parse_csv(content: bytes) -> tuple[list[str], list[dict[str, Any]], str | None]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    matrix = [[_number_or_text(value) for value in row] for row in csv.reader(io.StringIO(text), dialect)]
    return _rows_from_matrix(matrix, None)


def _parse_xlsx(content: bytes) -> tuple[list[str], list[dict[str, Any]], str | None]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"The Excel workbook could not be read: {exc}") from exc
    sheet = workbook["Everyday Activity"] if "Everyday Activity" in workbook.sheetnames else workbook.active
    matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
    return _rows_from_matrix(matrix, sheet.title)


def _parse_xls(content: bytes) -> tuple[list[str], list[dict[str, Any]], str | None]:
    try:
        workbook = xlrd.open_workbook(file_contents=content)
        sheet = workbook.sheet_by_index(0)
        matrix = [sheet.row_values(index) for index in range(sheet.nrows)]
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"The legacy Excel workbook could not be read: {exc}") from exc
    return _rows_from_matrix(matrix, sheet.name)


def _dataset_out(dataset: ImportedDataset) -> dict[str, Any]:
    return {
        "id": str(dataset.id), "name": dataset.name, "original_filename": dataset.original_filename,
        "file_type": dataset.file_type, "sheet_name": dataset.sheet_name, "row_count": dataset.row_count,
        "column_count": dataset.column_count, "columns": dataset.columns, "imported_at": dataset.imported_at,
    }


def _owned_dataset(dataset_id: str, current_user: User, db: Session) -> ImportedDataset:
    try:
        parsed_id = UUID(dataset_id)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail="Imported dataset not found") from exc
    dataset = db.query(ImportedDataset).filter(ImportedDataset.id == parsed_id, ImportedDataset.user_id == current_user.id).first()
    if not dataset:
        raise HTTPException(status_code=404, detail="Imported dataset not found")
    return dataset


@router.get("")
def list_datasets(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    datasets = db.query(ImportedDataset).filter(ImportedDataset.user_id == current_user.id).order_by(ImportedDataset.imported_at.desc()).all()
    return [_dataset_out(dataset) for dataset in datasets]


@router.post("/import", status_code=status.HTTP_201_CREATED)
async def import_dataset(
    request: Request,
    file: UploadFile = File(...),
    name: str | None = Form(None),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    filename = Path(file.filename or "dataset").name
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Upload a CSV, XLSX, or XLS file.")
    content = await file.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="Dataset files must be 10 MB or smaller.")
    if not content:
        raise HTTPException(status_code=422, detail="The uploaded file is empty.")

    columns, rows, sheet_name = _parse_csv(content) if extension == ".csv" else _parse_xlsx(content) if extension == ".xlsx" else _parse_xls(content)
    recognized = {alias for aliases in ALIASES.values() for alias in aliases}.intersection(columns)
    recognized.update(column for column in columns if column.endswith("_completed") or column.endswith("_goal_progress"))
    if not recognized:
        raise HTTPException(status_code=422, detail="No supported analytics columns were found. Include fields such as date, income_inr, total_expenses_inr, study_hours, habit_completion_rate, or overall_risk_score.")

    dataset = ImportedDataset(
        user_id=current_user.id,
        name=(name or Path(filename).stem).strip()[:180],
        original_filename=filename,
        file_type=extension.removeprefix(".").upper(),
        sheet_name=sheet_name,
        row_count=len(rows),
        column_count=len(columns),
        columns=columns,
        rows=rows,
    )
    db.add(dataset)
    db.commit()
    db.refresh(dataset)
    log_activity(
        db=db, user_id=current_user.id, action_type="DATASET_IMPORT", endpoint="/api/v1/datasets/import",
        ip_address=request.client.host if request.client else None, user_agent=request.headers.get("user-agent"),
        status_code=201, metadata={"dataset_id": str(dataset.id), "filename": filename, "rows": len(rows)},
    )
    result = _dataset_out(dataset)
    result["recognized_columns"] = sorted(recognized)
    return result


@router.get("/{dataset_id}/preview")
def preview_dataset(dataset_id: str, limit: int = 12, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    dataset = _owned_dataset(dataset_id, current_user, db)
    return {**_dataset_out(dataset), "rows": dataset.rows[:max(1, min(limit, 50))]}


@router.get("/{dataset_id}/dashboard-analysis")
def dataset_dashboard_analysis(dataset_id: str, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    dataset = _owned_dataset(dataset_id, current_user, db)
    return {"dataset": _dataset_out(dataset), **dashboard_analysis(dataset.rows)}


@router.get("/{dataset_id}/forecast-summary")
def dataset_forecast_summary(dataset_id: str, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    dataset = _owned_dataset(dataset_id, current_user, db)
    return {"dataset": _dataset_out(dataset), **forecast_analysis(dataset.rows)}


@router.delete("/{dataset_id}")
def delete_dataset(dataset_id: str, request: Request, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    dataset = _owned_dataset(dataset_id, current_user, db)
    filename = dataset.original_filename
    db.delete(dataset)
    db.commit()
    log_activity(
        db=db, user_id=current_user.id, action_type="DATASET_DELETE", endpoint=f"/api/v1/datasets/{dataset_id}",
        ip_address=request.client.host if request.client else None, user_agent=request.headers.get("user-agent"),
        status_code=200, metadata={"filename": filename},
    )
    return {"message": "Dataset removed"}
